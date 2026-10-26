from pathlib import Path
import json
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
IMG = ROOT / 'images'
OUT = Path(r'D:\Downloads\Final Year Project\output\documents')
OUT.mkdir(parents=True, exist_ok=True)
DEST = OUT / '2026-10-26_FastWAM_RealRobot_Task_Recommendations.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.header_distance = sec.footer_distance = Cm(.8)

def font(run, size=10.5, bold=False, color='000000'):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return run

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal=doc.styles['Normal']
normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.18
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for name,size in [('Title',23),('Heading 1',17),('Heading 2',12)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(9 if name!='Title' else 0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(8.5)
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.styles['Caption'].paragraph_format.line_spacing=1.08

def p(text='', boldlead=None, size=None, after=None, keep=False):
    a=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        font(a.add_run(boldlead),size or 10.5,True)
        font(a.add_run(text[len(boldlead):]),size or 10.5)
    else: font(a.add_run(text),size or 10.5)
    a.paragraph_format.keep_with_next=keep
    if after is not None: a.paragraph_format.space_after=Pt(after)
    return a

def h(text, new=False, level=1):
    a=doc.add_paragraph(text,'Heading '+str(level))
    if new: a.paragraph_format.page_break_before=True
    return a

def link(par, label, url, size=9):
    hl=OxmlElement('w:hyperlink')
    hl.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Calibri');rf.set(qn('w:eastAsia'),'Microsoft YaHei');rp.append(rf)
    c=OxmlElement('w:color');c.set(qn('w:val'),'225577');rp.append(c)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));rp.append(sz)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u)
    r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);hl.append(r);par._p.append(hl)

def source(n, locator):
    urls={1:'2607.28405v1',2:'2609.04193v1',3:'2609.20709v1',4:'2610.10079v1'}
    a=p(f'来源 [{n}] {locator}。',size=8.7,after=6)
    link(a,' 打开原文','https://arxiv.org/html/'+urls[n],8.7)

assets=[]
def picture(filename, caption, width=16.6, crop=None, stem=None):
    path=IMG/filename
    if crop:
        with Image.open(path) as im:
            assert 0<=crop[0]<crop[2]<=im.width and 0<=crop[1]<crop[3]<=im.height
            path=IMG/((stem or Path(filename).stem+'_crop')+'.png')
            im.crop(crop).save(path)
    a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER
    a.paragraph_format.space_after=Pt(3);a.paragraph_format.keep_with_next=True
    shape=a.add_run().add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set('descr',caption)
    c=doc.add_paragraph(caption,'Caption')
    assets.append({'asset':filename,'crop':crop,'caption':caption})
    return c

def table(headers, rows, widths, size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):c.text=txt
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for tag,v in [('top',65),('bottom',65),('left',80),('right',80)]:
                el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            for a in cell.paragraphs:
                a.paragraph_format.space_after=Pt(1);a.paragraph_format.line_spacing=1.08
                for run in a.runs:font(run,size,ri==0)
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    p('',after=1,size=2)
    return t

header=sec.header.paragraphs[0]
font(header.add_run('Fast WAM 真机任务参考'),8)
header.paragraph_format.space_after=Pt(0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run('文献核对 2026年10月26日  ·  '),8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

doc.core_properties.title='Fast WAM 单臂真机任务推荐'
doc.core_properties.subject='面向 BF16 与 Microscaling VQ inference engine 对照的真实机器人任务选型'
doc.core_properties.author='Research notes'

# Page 1
doc.add_paragraph('Fast WAM 单臂真机任务推荐','Title')
p('面向 BF16 与 Microscaling VQ inference engine 的成功率和加速比对照',size=12,after=12)
p('建议先以“物体放入篮子＋Stack bowls”建立静态对照，再从 Hang mug 或颜色指令放置中增加一项。前者有原生 Fast-WAM 单臂结果，后者有 Fast-WAM-IDM 单臂结果。若重点验证推理延迟对闭环表现的影响，再加入传送带 Carrot to plate。')
p('本报告按“单机械臂＋夹爪”的计划筛选，共列出七项主推荐任务。所有成功率均为文献结果；任务优先级、落地步骤及建议成功判据是为本项目提出的设计建议，尚未在我们的硬件上验证。')
h('候选任务总览',level=2)
table(['任务','已有模型证据','文献成功率','建议顺序'],[
    ['物体放入篮子','QuantWAMs 中的 Fast-WAM；单臂任务','FP16 8/10\nW4A4 8/10','首轮基础任务'],
    ['Stack bowls','MoWAM 中的 Fast-WAM；FR3 单臂','85%\nn = 20','首轮精度任务'],
    ['Hang mug','RealtimeWAM 中的原生 FastWAM；FR3','72%\nn = 50','静态进阶'],
    ['颜色指令放置','GIFT 中的 Fast-WAM-IDM；xArm7','7/10','采用 IDM 时优先'],
    ['Close drawer','MoWAM 中的 Fast-WAM；FR3 单臂','20%\nn = 20','BF16 先验证可做性'],
    ['指定货架层放置','GIFT 中的 Fast-WAM-IDM；xArm7','4/10','语义与三维定位进阶'],
    ['Carrot to plate','RealtimeWAM 中的原生 FastWAM；FR3','46%\nn = 50','动态扩展']
],[3.5,7.0,3.0,3.9],9.1)
p('如何读这张表  不同论文使用不同机器人、数据和执行设置，不能把各行成功率直接当作统一难度排行榜。FP16 与 BF16 也不是同一精度标签。QuantWAMs 的 8/10 对照提供量化实机先例，但小样本不支持“精度等价”结论。',boldlead='如何读这张表',size=9.5)
p('证据范围  QuantWAMs 的平台本身是双臂 AgiBot G2，苹果任务明确使用单臂。GIFT 的数字属于 Fast-WAM-IDM。MoWAM 和 RealtimeWAM 的原生 Fast-WAM 行是对照模型，图片中的作者方法演示不自动等于该对照行的 rollout。',boldlead='证据范围',size=9.5)
source(1,'§4.5，Table 6');source(2,'§5，Table 8');source(3,'§IV，Table II');source(4,'§4.5，Table 4')

# Page 2
h('硬件与相机参考',new=True)
p('任务可以借鉴，机器人接口与 checkpoint 仍需适配。选择已有单臂实测的平台布置，有助于减少任务设计的不确定性；它不意味着原 checkpoint 可以直接控制另一台机械臂。')
table(['来源','实机平台','相机披露','数据与测试规模'],[
    ['QuantWAMs [1]','AgiBot G2 上半身\n双臂各 7 DoF＋夹爪','头部与腕部三相机系统；苹果是单臂任务','苹果训练 200 条轨迹；每方法 10 次测试'],
    ['GIFT [2]','UFACTORY xArm7\n平行夹爪','2 × RealSense D435\n第三视角＋腕部','每项单臂任务 100 demos；每方法每任务 10 次测试'],
    ['MoWAM [3]','Franka Research 3\n单臂','未明确给出相机数量与型号','每任务 100 demos；20 次测试'],
    ['RealtimeWAM [4]','Franka Research 3\n末端型号未详述','RealSense D435i；未详述安装标定','未披露实机 demos 数量；每配置每任务 50 次测试']
],[2.4,4.3,5.1,5.6],8.8)
picture('gift_setup.png','图 1  GIFT Figure 8(a) 的单臂实机布置，PDF 第 14 页。左侧子图裁剪，保留 xArm7、第三视角与腕部相机；原图右侧双臂平台未纳入。[2]',width=15.1,crop=(0,0,1098,624),stem='gift_single_setup')
p('相机建议  现阶段可优先规划“固定第三视角 RGB＋腕部 RGB”，这与 GIFT 的单臂布置一致，方便兼顾全局目标和抓取细节。这是选型建议；最终视角数量、图像尺寸和拼接次序必须与选定的 Fast-WAM／IDM 训练代码及 checkpoint 输入一致。RealSense 的品牌不代表模型一定使用 depth。',boldlead='相机建议',size=9.7)
source(1,'§4.5，Figure 6');source(2,'§5-A1，Figure 8(a)');source(3,'§IV-A1');source(4,'Appendix A.1，Figure 8')

# Page 3
h('任务一 物体放入篮子',new=True)
p('推荐级别  首轮必选候选。它既有 Fast-WAM 实机记录，也有量化前后的同任务记录，最贴近本项目的验证目标。',boldlead='推荐级别')
picture('quantwams_real_0_0.png','图 2  QuantWAMs Figure 6 的平台与苹果任务局部，PDF 第 11 页。左侧是 AgiBot G2，右侧是苹果放篮示例；平台为双臂，所推荐的任务为单臂操作。[1]',width=9.3,crop=(0,0,407,312),stem='quantwams_apple')
h('原文已有结果',level=2)
p('QuantWAMs 在苹果放篮任务上使用 200 条训练轨迹；最终每种方法测试 10 次。Fast-WAM FP16 为 8/10，Atom* W4A4 为 5/10，加入 QuantWAMs 后为 8/10。该论文的 1.4×加速仅针对目标 WAM blocks，不能写成完整机器人控制周期的加速。[1]')
h('我们可以怎样做',level=2)
p('用一个篮子和一个易夹持的目标物体，在桌面标出若干起始位置。先固定目标类别和篮子位置，形成稳定 BF16 基线；之后再扩展起始位置和物体朝向。苹果可用尺寸适合夹爪的仿真水果替代，但报告应写明这是改编任务，不能称为原论文的精确复现。')
p('建议成功判据  目标物体被送入指定篮子，夹爪完全松开后在篮内稳定至少 3 秒；抓错物体、物体落在篮外或超过预设时限均为失败。时限和起始位置集合需在正式测试前固定。该判据是本项目建议，原文未提供这些数值容差。',boldlead='建议成功判据')
p('为何适合量化对照  抓取、抬升、运输、释放均可观察，失败类型直观。若量化模型丢失抓取精度，成功率容易反映；若仅推理更快，静态环境下未必增加成功率，因此同时报告 request latency。')
p('近似替代任务  MoWAM 的 Pick Banana 同为单臂取物，图中包含篮子。该论文原生 Fast-WAM 仅为 30%（n = 20），所以不建议仅因外观简单就把它视作低难度基线。[3]',size=9.6)
source(1,'§4.5，Figure 6，Table 6，PDF 第 11 页');source(3,'Figure 3，Table II')

# Page 4
h('任务二 颜色指令放置',new=True)
p('推荐级别  若选用 Fast-WAM-IDM，这是最值得优先加入的语义任务；如果最终使用原生 Fast-WAM，需另行确认该任务上的 BF16 可做性。',boldlead='推荐级别')
picture('gift_singlearm.png','图 3  GIFT Figure 9 的 Task 1 实机序列，PDF 第 15 页。作者展示颜色指令放置的操作过程；裁剪为原图第一行，不能据此将演示视为 Fast-WAM-IDM 对照模型的某次 rollout。[2]',crop=(0,0,1848,327),stem='gift_color_row')
h('原文已有设置与结果',level=2)
p('任务要求根据指令选定物体，并放入指定颜色的篮子，例如从多个物体中选出黄椒并放入蓝色篮子。GIFT 使用 xArm7、平行夹爪以及第三视角和腕部两台 D435；每任务 100 条 teleoperation demos。Table 8 的 Fast-WAM-IDM 为 7/10，GIFT-WAM-IDM 为 9/10。[2]')
h('我们可以怎样做',level=2)
p('桌上放两个颜色明显不同的容器，以及两个外形容易区分、都适合夹持的物体。指令同时指定“抓什么”和“放哪里”。先避免太多遮挡，保持两个篮子始终在视野内；抓取稳定后，再增加位置变化。这样既保留原任务的语义选择，又不会让早期结果主要受抓取硬件限制。')
p('每条 demonstration 应绑定实际执行的语言指令，并保存同步图像、机器人状态和动作。训练集应覆盖各目标物体与目标颜色组合，测试使用独立轨迹和预先定义的布局。这里是采集设计建议，文献中的 100 demos 不代表换机器人后必然足够。')
h('建议成功判据',level=2)
p('只有“正确物体进入正确颜色篮子，并在释放后稳定至少 3 秒”计为成功。物体选错、篮子选错、运输中掉落和超时分别记录，但主表统一作为二元失败。成功标签由整条任务决定，不能把“抓到任意物体”算作完成。')
h('对本项目的价值',level=2)
p('它能区分“仍然会抓取”和“仍然按语言选对目标”。在 BF16 与量化两臂中保持指令文本、颜色布局、相机曝光和 action chunk 设置一致，才能把结果与模型执行方式联系起来。')
source(2,'§5-A，§5-B，Table 8，Figure 9；PDF 第 14–15 页')

# Page 5
h('任务三 Stack bowls',new=True)
p('推荐级别  首轮精度任务。它在已有 Fast-WAM 单臂实验中具有较高成功率，又比宽松篮筐取放更依赖末端定位，适合检测量化后细微动作误差。',boldlead='推荐级别')
picture('mowam_tasks.png','图 4  MoWAM Figure 3 的 Stack Bowls 实机序列，PDF 第 5 页。裁剪为中间一行，展示一个碗被放到另一个碗上／内的过程；图片是作者任务演示。[3]',crop=(0,247,1230,489),stem='mowam_bowls')
h('原文已有设置与结果',level=2)
p('MoWAM 在 Franka Research 3 上测试 Pick Banana、Stack Bowls 和 Close Drawer，每任务采集 100 条 demos，并进行 20 次 evaluation trials。Table II 的原生 Fast-WAM 在 Stack Bowls 上为 85%，MoWAM 为 95%。相机数量和型号在该实机设置中未明确披露。[3]')
h('我们可以怎样做',level=2)
p('准备两个尺寸与夹爪匹配、容易稳定叠放或套放的碗，先固定底碗位置，并在几个桌面标记间改变上碗的初始位置。演示应覆盖接近、夹持、抬起、对齐、缓慢下降和释放全过程。正式论文应说明这里的 stack 是叠放还是套放，以及两个碗的尺寸关系。')
p('第一版不要同时引入易变形碗、极薄边缘和拥挤遮挡，否则失败可能主要来自夹持条件。任务稳定后，可把底碗位置变化作为独立测试条件，但不要把增加难度后的测试与原静态条件混成一个成功率。')
h('建议成功判据',level=2)
p('上碗放到指定底碗上／内，释放后保持预定的稳定叠放状态至少 3 秒，且不依赖夹爪支撑。底碗被推走、上碗倾覆、未完成释放或超时均为失败。建议在测试前拍摄一张合格终态和一张失败终态作为判分参考。')
h('对本项目的价值',level=2)
p('它将“能否完成大范围运输”进一步收紧为“能否在接触前对齐并稳定释放”。与任务一同时使用，可以观察量化退化是否主要出现在精细放置阶段。原文 85% 只是该平台和该协议的结果，不能作为我们新硬件的预期成功率。')
source(3,'§IV-A1，Figure 3，Table II；PDF 第 5–6 页')

# Page 6
h('任务四 Close drawer',new=True)
p('推荐级别  条件式候选。它能补充接触交互，但原文 Fast-WAM 成功率较低，应先确认我们的 BF16 能稳定执行，再决定是否放入主实验。',boldlead='推荐级别')
picture('mowam_tasks.png','图 5  MoWAM Figure 3 的 Close Drawer 实机序列，PDF 第 5 页。裁剪为下方一行。推荐的是关抽屉，图示不应被解释成拉开抽屉任务。[3]',crop=(0,496,1230,736),stem='mowam_drawer')
h('原文已有设置与结果',level=2)
p('同属 MoWAM 的 FR3 单臂实机任务，每任务 100 条 demos、20 次测试。Table II 的 Fast-WAM 为 20%，MoWAM 为 80%。这说明论文中该任务对原生 Fast-WAM 并不轻松；不能因为它看起来只有一次推动就认定是简单任务。[3]')
h('我们可以怎样做',level=2)
p('使用固定在桌面的轻型抽屉箱，选择平滑、行程可重复的滑轨。将抽屉打开到事先标定的若干位置，让末端接近合适的接触面，再把抽屉推到关闭状态。箱体要固定，避免模型通过推动整个箱子达到看似接近的图像终态。')
p('先控制抽屉初始开度和接触位置；任务稳定后再扩展这些条件。演示应包含接触前减速、持续推动和完成后撤离，尽量减少人为干预。若实际末端需要先抓住把手再推回，应将这一实现写清楚，不能与单纯推动的任务混称为精确复现。')
h('建议成功判据',level=2)
p('夹爪撤离后抽屉保持关闭，例如前沿与关闭基准线的间隙不超过 1 cm，箱体不发生明显移位，并在时限内完成。1 cm 是建议起点，应按抽屉结构固定最终容差；它不是原论文披露的标准。')
h('何时纳入主结果',level=2)
p('如果 BF16 已能稳定关上抽屉，量化对照可以检验接触阶段的动作稳定性；如果 BF16 成功率接近下限，两个方法都失败会掩盖量化影响。此时应先完善数据与控制接口，或把它保留为附加压力测试，主表采用其他可完成的静态任务。')
source(3,'§IV-A1，Figure 3，Table II；PDF 第 5–6 页')

# Page 7
h('任务五 指定货架层放置',new=True)
p('推荐级别  Fast-WAM-IDM 的进阶候选。它同时考察物体选择和三维目标定位，适合在桌面颜色放置稳定之后加入。',boldlead='推荐级别')
picture('gift_singlearm.png','图 6  GIFT Figure 9 的 Task 2 实机序列，PDF 第 15 页。裁剪为原图第二行；根据语言将指定物体放到指定货架层。[2]',crop=(0,333,1848,649),stem='gift_shelf_row')
h('原文已有设置与结果',level=2)
p('GIFT 的单臂 Task 2 是 layer-conditioned placement，例如将指定橙子放到货架最上层。机器人与相机设置同任务二，每任务 100 条 demos。Fast-WAM-IDM 在原始条件下为 4/10，GIFT-WAM-IDM 为 8/10；原文指出目标可用区域较窄。[2]')
h('我们可以怎样做',level=2)
p('准备一个固定在桌面的两层或三层小货架，目标物体尺寸明显小于层高与入口宽度。让语言明确指向“上层／下层”等预定义位置。第一轮应避免把“必须绕过横梁”“必须从特定侧面进入”等额外约束叠加进去。')
p('采集时覆盖物体起始位置和不同目标层，保留机械臂、夹爪和货架边缘的可见性。动作执行前需确认目标层在机械臂工作空间内，腕部相机不会在接近时完全遮住入口；这些属于任务可做性的设置，不应由量化对照承担。')
h('建议成功判据',level=2)
p('正确物体完整放置在指定层的预先标记区域内，释放后稳定至少 3 秒；放错层、跨出区域、落下或碰撞后卡住均记为失败。区域可以用可移除标线定义，但训练与测试的判分边界必须保持一致。')
h('扩展时需要注意',level=2)
p('GIFT 还在该任务测试旋转彩灯及更换桌布；Fast-WAM-IDM 的两级结果为 3/10 和 2/10。这些是特定扰动条件的小样本结果。我们若加入光照或背景变化，应单列条件，并先完成原始条件下的 BF16 与量化比较。')
source(2,'§5，Table 8–9，Figure 9；PDF 第 14–15 页')

# Page 8
h('任务六 Hang mug',new=True)
p('推荐级别  静态进阶优先候选。它有明确的原生 FastWAM 单臂实测，能够以很少的道具形成比取放更强的姿态和对齐约束。',boldlead='推荐级别')
picture('realtimewam_real_0_0.png','图 7  RealtimeWAM Figure 7 的 C 行 Hang mug，PDF 第 9 页。裁剪为该任务的实机序列，展示杯柄与挂架对齐及悬挂。[4]',crop=(0,202,816,296),stem='realtime_mug')
h('原文已有设置与结果',level=2)
p('RealtimeWAM 使用 FR3 和 RealSense D435i，Hang mug 在水平桌面上进行。每任务、backbone 和 inference configuration 各做 50 次测试。原生 FastWAM 为 72%，RealtimeWAM 加速版为 70%；对应均值推理延迟为 211.08 ms 和 25.07 ms。这是执行加速方法的对照，不是量化结果。[4]')
h('我们可以怎样做',level=2)
p('选择有足够大杯柄、适合夹爪抓持的轻质杯子，以及固定牢靠、入口清晰的单个挂钩。先固定挂钩位置与杯柄朝向，只改变杯子的少量起始位置。任务动作包括抓杯、调整姿态、让杯柄越过挂钩、下降到承重点、释放并撤离。')
p('杯柄尺寸与挂钩尺寸决定对齐难度。报告应给出二者大致尺寸和初始姿态范围，并保持 BF16 与量化两臂使用同一套道具。若杯子太重或夹持点不稳定，失败可能反映夹爪条件，难以解释为模型量化带来的退化。')
h('建议成功判据',level=2)
p('夹爪撤离后，杯子仅通过杯柄稳定挂在指定挂钩上至少 3 秒；杯底放在桌面、杯身被夹爪托住或杯子随后滑落均不计成功。将杯子接近挂架不能算作任务完成。')
h('为什么值得保留',level=2)
p('静态目标有利于先看成功率保持；杯柄对齐又对动作细节敏感。论文中大幅降延迟并未提高该任务成功率，提醒我们：静态任务的主要结论可以是“成功率保持且推理更快”，不必要求每个任务成功率都上升。')
source(4,'§4.5，Table 4，Figure 7；Appendix A，Table 5；PDF 第 9、14–15 页')

# Page 9
h('任务七 传送带 Carrot to plate',new=True)
p('推荐级别  动态扩展的第一候选。需要额外传送带，适合在静态基线完成后检验更短推理延迟能否改善移动目标操作。',boldlead='推荐级别')
picture('realtimewam_real_2_0.png','图 8  RealtimeWAM Figure 8，PDF 第 14 页。裁剪为原图上排正视图，从左到右为传送带、斜面桌和水平桌。本任务使用左侧传送带。[4]',crop=(0,0,660,209),stem='realtime_setup_front')
picture('realtimewam_real_0_0.png','图 9  RealtimeWAM Figure 7 的 B 行 Carrot to plate，PDF 第 9 页。裁剪为将传送带上的胡萝卜移至盘子的实机序列。[4]',crop=(0,100,816,195),stem='realtime_carrot')
h('原文已有设置与结果',level=2)
p('原文传送带速度为 1 m/min。原生 FastWAM 为 46%（n = 50），RealtimeWAM 为 74%（n = 50），均值推理延迟分别为 209.74 ms 和 21.86 ms。机器人以 20 Hz 执行，单次请求预测 32 个 future actions，采用同步推理与执行。[4]')
h('我们可以怎样做',level=2)
p('先保持盘子固定、传送带速度固定，以胡萝卜或形状相近的易夹持目标做单目标转移。正式测试使用成对的目标起始位置和启动时序。不要在 BF16 和量化两臂间分别调速度、replan 频率或可等待时间，否则速度与执行策略会同时变化。')
p('建议成功判据  在目标离开预定抓取区域前完成抓取，并把目标送入盘子；释放后在盘内稳定至少 3 秒。超过任务时限、目标离开抓取区域或掉出盘子均失败。抓取区域与时限是本项目需要预先固定的定义，并非原文公开的数值判据。',boldlead='建议成功判据')
p('进一步的动态候选  同文还测试 Tennis ball to bowl、Three shuttlecocks 和 Rolling can to bowl。原生 FastWAM 分别为 26%、24%、10%（各 n = 50）。Rolling can 需要一侧抬高 5 cm 的斜面且基线偏低，更适合作为后续压力测试。',boldlead='进一步的动态候选',size=9.5)
source(4,'Table 4–5，Figure 7–8，Appendix A；PDF 第 9、14–15 页')

# Page 10
h('建议最终选哪些任务',new=True)
p('如果第一轮只做三个任务，我建议“物体放入篮子＋Stack bowls＋Hang mug”。它们覆盖宽松取放、精细放置和姿态对齐，均有 Fast-WAM 的真实机器人证据，也都适合单臂＋夹爪。若采用 Fast-WAM-IDM 并希望突出语言选择，可以用颜色指令放置替换 Hang mug。')
table(['实验组合','任务','适用目的'],[
    ['最小可行组合','物体放入篮子＋Stack bowls','先验证量化模型能稳定执行，并观察定位精度保持'],
    ['推荐主实验组合','上述两项＋Hang mug\n或在 IDM 下改用颜色指令放置','增加对齐或语义选择能力，仍保持单臂静态布置'],
    ['延迟敏感扩展','主实验＋Carrot to plate','在移动目标中观察更短 request latency 的闭环价值'],
    ['后续压力测试','Close drawer、指定货架层、Rolling can','在 BF16 已可完成的前提下扩大难度覆盖']
],[3.2,6.6,7.6],9.4)
h('把主表限制为成功率和 speedup',level=2)
p('两臂采用同一个实机微调 checkpoint：一臂 BF16，另一臂为我们的量化模型＋C／decoder／buffer inference engine。保持相机、预处理、语言、动作反归一化、采样步数、CFG、action chunk、replan 频率及安全控制相同。量化校准使用训练集，最终测试场景不进入校准或调参。')
p('成功率  每任务报告成功数／总尝试数和百分比。两臂使用同一批预设布局、指令与复位条件，并交替或随机化执行顺序，减少物体磨损、光照漂移和人工复位习惯的影响。任务失败和人工接管都保留在分母内；基础设施故障的排除规则提前定义。',boldlead='成功率')
p('Speedup  采用完整 policy request 的平均延迟比：speedup = mean latency(BF16) / mean latency(quantized + engine)。计时从已准备好的观测进入策略接口，到可交付机器人的反归一化动作返回；包括输入处理、GPU 执行、解码、数据传输和必要同步。若解码或 buffer 准备发生在请求内，它们必须被计入。',boldlead='Speedup')
p('编译、初次加载和 warmup 单列；相机采集、网络和机械臂运动若不计入，应在表注说明。同一硬件、batch size 和 warmup 规则下比较，表注补充动作请求次数以及平均延迟毫秒值，读者才能判断 speedup 的测量范围。不要用单个 Linear 或 decoder kernel 的倍数代替策略端到端倍数。')
h('开始采集前需要确定的最小信息',level=2)
p('机械臂和夹爪型号；控制命令是关节量还是末端位姿及其单位；控制频率；相机视角与时间同步；数据存储格式；动作与图像的时间对齐。原论文使用某个任务，不代表示范数据在不同机器人间可以直接混用。先以小规模 BF16 pilot 检查接口和任务可做性，再冻结正式的任务集与评分标准。')
p('本报告是任务选型与文献参考，不是已执行的实验协议。正式试验次数、置信区间和可接受的成功率下降幅度应在采集／评测前另行确定，文献中的 10、20、50 次只用于解释其原始证据。',size=9.3)

# Page 11
h('原始来源与配图出处',new=True)
p('引用的是下列固定版本。页码均指该版本 PDF 的物理页序；正文同时保留章节、图号和表号，便于在 HTML 与 PDF 间交叉定位。')
refs=[
    (1,'Jiacheng Zhou et al.','QuantWAMs: Calibrating at the Right Granularity for World Action Models','2607.28405v1','2026-07-30','§4.5；Figure 6、Table 6；PDF 第 11 页。用于单臂苹果放篮、量化实机结果和 G2 平台设置。'),
    (2,'Yupeng Zheng et al.','GIFT: Guided Intermediate Feature Training via Action-Oriented Structural Supervision for Robotic Manipulation','2609.04193v1','2026-09-03','§5；Figure 8–9、Table 8–9；PDF 第 14–15 页。用于 xArm7 硬件、颜色／层位任务及 Fast-WAM-IDM 结果。'),
    (3,'Jiayu Wang et al.','MoWAM: Explicit Future Motion Prediction for Efficient World Action Models','2609.20709v1','2026-09-17','§IV-A1 与实机结果；Figure 3、Table II；PDF 第 5–6 页。用于 FR3 的香蕉、叠碗与关抽屉任务。'),
    (4,'Huanan Liu et al.','RealtimeWAM: How Fast Can I Run My World Action Model?','2610.10079v1','2026-10-07','§4.5；Appendix A、B.1–B.2、C.1；Figure 7–8、Table 4–5；PDF 第 9、14–16 页等。用于静态与动态任务、原生 FastWAM 结果及计时边界。')
]
for n,authors,title,aid,date,loc in refs:
    p(f'[{n}] {authors} {title}.',size=10.3,keep=True)
    p(f'arXiv:{aid}，版本日期 {date}。{loc}',size=9.2,after=2)
    a=doc.add_paragraph();a.paragraph_format.space_after=Pt(9)
    link(a,'HTML 原文','https://arxiv.org/html/'+aid,9)
    font(a.add_run('    '),9)
    link(a,'PDF 原文','https://arxiv.org/pdf/'+aid,9)
h('配图使用说明',level=2)
p('本报告共使用九幅实机图，均取自原作者论文的公开图片资源。仅做局部裁剪与等比例缩放，没有生成或补画场景。图注注明了论文图号、PDF 页码和来源编号；作者方法演示图用于说明任务过程，不作为其 Fast-WAM 对照行成功率的逐次证据。',size=9.7)
for label,url in [
    ('图 1  GIFT setup.png','https://arxiv.org/html/2609.04193v1/setup.png'),
    ('图 2  QuantWAMs Real-tasks.png','https://arxiv.org/html/2607.28405v1/Real-tasks.png'),
    ('图 3、6  GIFT singlearm.png','https://arxiv.org/html/2609.04193v1/singlearm.png'),
    ('图 4、5  MoWAM demo.png','https://arxiv.org/html/2609.20709v1/images/demo.png'),
    ('图 7、9  RealtimeWAM task sequences','https://arxiv.org/html/2610.10079v1/figure06_real_world.png'),
    ('图 8  RealtimeWAM real world setup','https://arxiv.org/html/2610.10079v1/appendix_real_world_setup.png')
]:
    a=doc.add_paragraph();a.paragraph_format.space_after=Pt(4);link(a,label,url,9)
p('正式论文、公开演示文稿或投稿材料若再次使用这些图片，应按原文许可与出版要求处理署名及转载授权。本文件用于任务选型参考。',size=9)
for style in doc.styles:
    for border in list(style._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(DEST)
(ROOT/'figure_manifest.json').write_text(json.dumps(assets,ensure_ascii=False,indent=2),encoding='utf-8')
assert DEST.exists() and len(doc.inline_shapes)==9
assert len(doc.tables)==3
print(str(DEST))
